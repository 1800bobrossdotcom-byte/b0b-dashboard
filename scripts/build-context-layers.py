#!/usr/bin/env python3
"""Build the GL map's context layers from source files the spider has ledgered.

Context layers are infrastructure drawn under the evidence for orientation. They
are not evidence and are never drawn as edges: a site near a power plant is not
connected to it. Each layer is built only from a source whose licence was read
at the source (research/MAP-CONTEXT-SOURCES.md, 22 Oct 2026), from the exact
bytes the ledger recorded - the script refuses a body whose sha256 differs.

Output goes to public/data-context/ (Vercel static output; site/ is near the
function size cap). Each file carries its own provenance block.

  python3 scripts/build-context-layers.py            all layers
  python3 scripts/build-context-layers.py flares     one layer

Layers:
  power-plants.json   WRI Global Power Plant Database v1.3.0, CC BY 4.0.
                      Frozen June 2021, unmaintained by WRI since early 2022.
  flares.json         World Bank GFMR flare volumes by flare, 2012-2025, under the
                      World Bank dataset terms (CC BY 4.0 unless labelled otherwise;
                      no endorsement may be implied).
  data-centres.json   Epoch AI "AI data centers", CC BY. Placed by
                      scripts/geocode-context.py (scripts/context-geocodes.json).
  oil-gas.json        OGIM v3.0 (EDF / MethaneSAT), CC BY 4.0: refineries, LNG
                      facilities, offshore platforms, petroleum terminals. Not its
                      flaring layer (licensed to EDF only) and not its 1.86 million
                      pipeline segments. The 3.4 GB GeoPackage lives outside the repo
                      (B0B_OGIM, default /home/user/.b0b-sources/ogim/OGIM_v3.0.gpkg).

Not built: submarine cables, landing points and IXPs - no openly licensed
source exists (TeleGeography sells its data; PeeringDB and IXPDB bar bulk
redistribution; PCH is non-commercial).
"""
import csv
import hashlib
import io
import json
import os
import sqlite3
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CACHE = os.path.join(ROOT, 'research', 'sources', 'cache')
OUT = os.path.join(ROOT, 'public', 'data-context')

GPPD_URL = 'https://raw.githubusercontent.com/wri/global-power-plant-database/master/output_database/global_power_plant_database.csv'
GPPD_SHA = '4b1f93e0fd93664f18684d9b05d0a52ed9658c6a8cf0d21ff2520791379ba7fc'
FLARE_URL = 'https://thedocs.worldbank.org/en/doc/b34e0c054bb3fe3695e70154c28eef3f-0400072026/related/Flare-Volume-Estimates-by-individual-Flare-Location-2012-2025.xlsx'
FLARE_SHA = '8b6052d23c09584d5964d6549026961b5ee8529b9b882edb4073fd6e9bc40354'
EPOCH_URL = 'https://epoch.ai/data/data_centers/data_centers.csv'
EPOCH_SHA = '271e08676cbf089705f32d86439ac5ae87cef61f849a76aa8ccc6b2542eb640a'
OGIM_URL = 'https://zenodo.org/records/22835235/files/OGIM_v3.0.gpkg?download=1'
OGIM_SHA = 'd324540fa13a29ce731aee9753077dffee81470abbe7652c9488b7238e80b04c'
OGIM_PATH = os.environ.get('B0B_OGIM', '/home/user/.b0b-sources/ogim/OGIM_v3.0.gpkg')   # absolute: run as root, ~ is /root


def cached(url, sha):
    p = os.path.join(CACHE, sha[:2], sha)
    if not os.path.exists(p):
        sys.exit('missing cache body for %s\n  run: python3 scripts/spider/spider.py fetch %s' % (url, url))
    data = open(p, 'rb').read()
    got = hashlib.sha256(data).hexdigest()
    if got != sha:
        sys.exit('cache body for %s hashes %s, ledger says %s - refusing' % (url, got, sha))
    return data


def sha_file(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for block in iter(lambda: f.read(1 << 22), b''):
            h.update(block)
    return h.hexdigest()


def point(lng, lat, props):
    return {'type': 'Feature', 'geometry': {'type': 'Point', 'coordinates': [round(lng, 4), round(lat, 4)]}, 'properties': props}


def clean(v):
    """Source placeholders ('N/A', 'NA', -999, 1900-01-01) become absent."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return None if v <= -999 else v
    v = str(v).strip()
    return None if v in ('', 'N/A', 'NA', 'None', '1900-01-01') else v


def power_plants():
    rows = list(csv.DictReader(cached(GPPD_URL, GPPD_SHA).decode('utf-8').splitlines()))
    feats = []
    for r in rows:
        try:
            lat, lng, mw = float(r['latitude']), float(r['longitude']), float(r['capacity_mw'])
        except ValueError:
            continue
        others = [r[k] for k in ('other_fuel1', 'other_fuel2', 'other_fuel3') if r[k]]
        p = {
            'name': r['name'],
            'fuel': r['primary_fuel'] + (' (also ' + ', '.join(others) + ')' if others else ''),
            'mw': round(mw, 1),
            'country': r['country_long'],
        }
        if r['commissioning_year']:
            p['year'] = str(int(float(r['commissioning_year'])))
        if r['source']:
            p['src'] = r['source']
        p['id'] = r['gppd_idnr']
        feats.append(point(lng, lat, p))
    # largest first, so the biggest plants draw underneath the small ones
    feats.sort(key=lambda f: -f['properties']['mw'])
    return feats, {
        'dataset': 'Global Power Plant Database v1.3.0 (released 2 June 2021)',
        'publisher': 'World Resources Institute and Data Contributors',
        'attribution': 'Global Energy Observatory, Google, KTH Royal Institute of Technology in Stockholm, Enipedia, World Resources Institute. 2021. Global Power Plant Database version 1.3.0. Accessed through https://datasets.wri.org/datasets/global-power-plant-database on 22 October 2026',
        'licence': 'CC BY 4.0 (https://creativecommons.org/licenses/by/4.0/)',
        'changes': 'Simplified by b0b.dev: fields reduced to name, fuel, capacity, country, commissioning year, per-row source and id; owner and generation columns dropped; coordinates rounded to 4 decimal places.',
        'caveats': 'Not maintained by WRI since early 2022. Plants of 1 MW and over, about 72% of world capacity by the publisher\'s estimate; the publisher notes that matching can merge two plants or split one.',
        'source_url': GPPD_URL, 'source_sha256': GPPD_SHA,
    }


def flares():
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(cached(FLARE_URL, FLARE_SHA)), read_only=True, data_only=True)
    ws = wb['Dashboard_file']
    rows = ws.iter_rows(values_only=True)
    head = list(next(rows))
    ix = {k: head.index(k) for k in head}
    by = {}
    for r in rows:
        fid = r[ix['id']]
        if fid is None:
            continue
        rec = by.setdefault(fid, {'years': {}, 'row': r})
        if r[ix['Year']] is not None:
            rec['years'][int(r[ix['Year']])] = r[ix['Flaring Vol (million m3)']] or 0
            if int(r[ix['Year']]) >= int(rec['row'][ix['Year']] or 0):
                rec['row'] = r
    feats = []
    for fid, rec in by.items():
        r, ys = rec['row'], rec['years']
        try:
            lat, lng = float(r[ix['lat']]), float(r[ix['lon']])
        except (TypeError, ValueError):
            continue
        active = sorted(y for y, v in ys.items() if v and v > 0)
        peak = max(ys.items(), key=lambda kv: kv[1] or 0) if ys else (None, 0)
        p = {
            'name': clean(r[ix['Field name']]) or 'Flare ' + str(fid),
            'country': clean(r[ix['Country']]),
            'v2025': round(ys.get(2025) or 0, 2),
            'peak': '%s (%.1f)' % (peak[0], peak[1] or 0) if peak[0] else None,
            'years': '%d of 14 (%s-%s)' % (len(active), active[0], active[-1]) if active else 'none',
            'field': clean(r[ix['Field type']]),
            'where': clean(r[ix['Location']]),
            'operator': clean(r[ix['Operator']]),
            'id': fid,
        }
        feats.append(point(lng, lat, {k: v for k, v in p.items() if v is not None}))
    feats.sort(key=lambda f: -f['properties']['v2025'])
    return feats, {
        'dataset': '2012 - 2025 Flare Volume Estimates by Individual Flare Location (Global Gas Flaring Tracker, June 2026)',
        'publisher': 'World Bank Global Flaring and Methane Reduction Partnership (GFMR)',
        'attribution': 'The World Bank: 2012 - 2025 Flare Volume Estimates by Individual Flare Location: NOAA, the Payne Institute at the Colorado School of Mines, World Bank/GFMR',
        'licence': 'World Bank dataset terms: CC BY 4.0 unless specifically labelled otherwise (this file is not labelled otherwise), with the World Bank\'s additional terms. The World Bank does not endorse this map or its use of the data.',
        'changes': 'Reduced by b0b.dev to one point per flare id: its latest location and metadata, the 2025 volume, the peak year and the years with a detected volume. Volumes are million cubic metres per year.',
        'caveats': 'Annual volumes estimated from satellite-detected heat (VIIRS Nightfire), not live detections. The publisher notes the field metadata may be out of date.',
        'source_url': FLARE_URL, 'source_sha256': FLARE_SHA,
    }


def data_centres():
    rows = list(csv.DictReader(io.StringIO(cached(EPOCH_URL, EPOCH_SHA).decode('utf-8'))))
    geo = json.load(open(os.path.join(ROOT, 'scripts', 'context-geocodes.json'), encoding='utf-8'))
    tag = lambda s: ', '.join(x.strip() for x in (s or '').split(',') if x.strip())   # keep Epoch's #confident/#likely tags
    feats, unplaced = [], []
    for r in rows:
        g = geo.get(r['Name']) or {}
        if g.get('lat') is None:
            unplaced.append(r['Name'])
            continue
        mw = clean(r['Current power (MW)'])
        p = {
            'name': r['Name'],
            'owner': tag(r['Owner']) or None,
            'users': tag(r['Users']) or None,
            'mw': float(mw) if mw else None,
            'h100e': format(round(float(clean(r['Current H100 equivalents']))), ',') if clean(r['Current H100 equivalents']) else None,
            'project': clean(r['Project']),
            'country': clean(r['Country']),
            'placed': ('at the street address (US Census geocoder)' if g['precision'] == 'address'
                       else 'at %s, the place its %s names' % (g['place'], g.get('named_in', 'record'))),
        }
        if 'anthropic' in (r['Users'] + r['Owner']).lower():
            p['disclosure'] = 'Anthropic, named here as a user, makes the AI model that helps draft this site.'
        feats.append(point(g['lng'], g['lat'], {k: v for k, v in p.items() if v is not None}))
    feats.sort(key=lambda f: -(f['properties'].get('mw') or 0))
    return feats, {
        'dataset': 'Epoch AI, "AI data centers" (updated 20 Oct 2026)',
        'publisher': 'Epoch AI',
        'attribution': "Epoch AI, 'AI data centers'. Published online at epoch.ai. Retrieved from 'https://epoch.ai/data/data-centers-documentation' [online resource]. Accessed 22 Oct 2026.",
        'licence': 'Creative Commons Attribution (CC BY)',
        'changes': 'Placed by b0b.dev: US street addresses by the US Census Bureau geocoder; other sites at the town their record names, from that town\'s Wikipedia coordinates. Fields reduced.',
        'caveats': 'About 95 of the largest AI sites, selected by compute; not a register of data centres. Owner and user tags (#confident / #likely / #speculative) are Epoch\'s own.',
        'unplaced': unplaced,
        'source_url': EPOCH_URL, 'source_sha256': EPOCH_SHA,
    }


OGIM_LAYERS = [('Crude_Oil_Refineries', 'refinery'), ('LNG_Facilities', 'LNG facility'),
               ('Offshore_Platforms', 'offshore platform'), ('Petroleum_Terminals', 'petroleum terminal')]


def oil_gas():
    if not os.path.exists(OGIM_PATH):
        sys.exit('missing %s - download %s (3.4 GB) there, or set B0B_OGIM' % (OGIM_PATH, OGIM_URL))
    got = sha_file(OGIM_PATH)
    if got != OGIM_SHA:
        sys.exit('OGIM file hashes %s, ledger says %s - refusing' % (got, OGIM_SHA))
    c = sqlite3.connect('file:%s?mode=ro' % OGIM_PATH, uri=True)
    feats = []
    for table, kind in OGIM_LAYERS:
        cols = [r[1] for r in c.execute('pragma table_info("%s")' % table)]
        for row in c.execute('select * from "%s"' % table):
            r = dict(zip(cols, row))
            if r['LATITUDE'] is None or r['LONGITUDE'] is None:
                continue
            cap = []
            if clean(r.get('LIQ_CAPACITY_BPD')):
                cap.append('%s barrels/day' % format(int(r['LIQ_CAPACITY_BPD']), ','))
            if clean(r.get('GAS_CAPACITY_MMCFD')):
                cap.append('%s million cubic feet/day' % format(round(r['GAS_CAPACITY_MMCFD'], 1), ','))
            p = {
                'name': (clean(r['FAC_NAME']) or kind).title() if clean(r['FAC_NAME']) else kind.capitalize(),
                'kind': kind,
                'type': clean(r['FAC_TYPE']),
                'status': clean(r['FAC_STATUS']),
                'operator': clean(r['OPERATOR']),
                'capacity': '; '.join(cap) or None,
                'country': (clean(r['COUNTRY']) or '').title() or None,
                'srcdate': clean(r['SRC_DATE']),
                'id': r['OGIM_ID'],
            }
            feats.append(point(r['LONGITUDE'], r['LATITUDE'], {k: v for k, v in p.items() if v is not None}))
    return feats, {
        'dataset': 'Oil and Gas Infrastructure Mapping (OGIM) database v3.0 (18 Sept 2026)',
        'publisher': 'Environmental Defense Fund / MethaneSAT',
        'attribution': 'Omara et al. 2023, Earth System Science Data 15, 3761, https://doi.org/10.5194/essd-15-3761-2023; OGIM v3.0, Zenodo record 22835235.',
        'licence': 'CC BY 4.0 (the Zenodo record\'s rights field)',
        'changes': 'Four point layers only (refineries, LNG facilities, offshore platforms, petroleum terminals); fields reduced; coordinates rounded to 4 decimal places. The flaring layer (licensed to EDF only) and the pipelines are not included.',
        'caveats': 'Compiled by EDF from public sources of variable spatial quality and date (each feature carries its source date). The open version omits about 300 Russian compressor stations (not among these layers).',
        'source_url': OGIM_URL, 'source_sha256': OGIM_SHA,
    }


LAYERS = {'power-plants': power_plants, 'flares': flares, 'data-centres': data_centres, 'oil-gas': oil_gas}


def main():
    os.makedirs(OUT, exist_ok=True)
    pick = [a for a in sys.argv[1:] if a in LAYERS] or list(LAYERS)
    for name in pick:
        feats, prov = LAYERS[name]()
        prov['features'] = len(feats)
        body = json.dumps({'type': 'FeatureCollection', 'provenance': prov, 'features': feats},
                          ensure_ascii=False, separators=(',', ':'))
        with open(os.path.join(OUT, name + '.json'), 'w', encoding='utf-8') as f:
            f.write(body)
        print('%-18s %6d features  %9d bytes' % (name + '.json', len(feats), len(body.encode('utf-8'))))


if __name__ == '__main__':
    main()
